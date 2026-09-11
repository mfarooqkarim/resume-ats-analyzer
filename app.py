import io
import json
import re
from typing import Any, Dict, List

import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document


# =========================================================
# Configuration
# =========================================================

st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)

MODEL_NAME = "gemini-2.5-flash"
MAX_FILE_SIZE_MB = 10
MAX_RESUME_CHARS = 30000
MAX_JOB_DESCRIPTION_CHARS = 15000


# =========================================================
# File extraction
# =========================================================

def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract text from a text-based PDF."""
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []

    for page in reader.pages:
        pages.append(page.extract_text() or "")

    return "\n".join(pages).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract text from paragraphs and tables in a DOCX file."""
    document = Document(io.BytesIO(file_bytes))
    content = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            content.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                content.append(" | ".join(row_text))

    return "\n".join(content).strip()


def extract_resume_text(uploaded_file) -> str:
    """Extract resume text based on file extension."""
    file_bytes = uploaded_file.getvalue()
    file_name = uploaded_file.name.lower()

    max_bytes = MAX_FILE_SIZE_MB * 1024 * 1024
    if len(file_bytes) > max_bytes:
        raise ValueError(
            f"File is too large. Please upload a file smaller than "
            f"{MAX_FILE_SIZE_MB} MB."
        )

    if file_name.endswith(".pdf"):
        text = extract_pdf_text(file_bytes)
    elif file_name.endswith(".docx"):
        text = extract_docx_text(file_bytes)
    else:
        raise ValueError("Only PDF and DOCX files are supported.")

    if not text.strip():
        raise ValueError(
            "No readable text was found. This may be a scanned or image-only "
            "PDF. Please upload a text-based PDF or DOCX file."
        )

    return text


# =========================================================
# Rule-based ATS scoring
# =========================================================

def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text.lower()).strip()


def extract_keywords(text: str) -> set:
    """Extract simple keyword-like terms from text."""
    stop_words = {
        "and", "the", "with", "for", "from", "that", "this", "your",
        "have", "has", "are", "was", "were", "will", "you", "our",
        "their", "they", "using", "into", "about", "over", "under",
        "years", "year", "work", "working", "team", "job", "role",
        "experience", "required", "preferred", "ability", "skills",
        "responsibilities", "description", "candidate", "strong",
        "knowledge", "including", "such", "other", "should", "must",
        "would", "could", "also", "more", "than", "some", "who",
        "what", "when", "where", "how", "our", "their", "its",
    }

    words = re.findall(
        r"[a-zA-Z][a-zA-Z0-9+#./-]{1,}",
        text.lower(),
    )

    return {
        word
        for word in words
        if word not in stop_words and len(word) >= 3
    }


def calculate_basic_ats_score(
    resume_text: str,
    job_description: str,
) -> Dict[str, Any]:
    """
    Calculate an approximate, transparent ATS compatibility score.

    This is not the official score of any ATS platform.
    """
    normalized_resume = normalize_text(resume_text)

    score = 0
    breakdown = []

    # 1. Content length: 15 points
    word_count = len(resume_text.split())

    if 250 <= word_count <= 1200:
        length_score = 15
    elif 150 <= word_count <= 1500:
        length_score = 10
    else:
        length_score = 5

    score += length_score
    breakdown.append({
        "category": "Resume content length",
        "score": length_score,
        "max_score": 15,
    })

    # 2. Standard sections: 20 points
    sections = {
        "contact information": ["email", "phone", "linkedin", "github"],
        "summary or objective": ["summary", "objective", "profile"],
        "experience": ["experience", "employment", "work history"],
        "education": ["education", "academic"],
        "skills": ["skills", "technical skills", "core competencies"],
    }

    section_score = 0
    detected_sections = []

    for section_name, keywords in sections.items():
        if any(keyword in normalized_resume for keyword in keywords):
            section_score += 4
            detected_sections.append(section_name)

    score += section_score
    breakdown.append({
        "category": "Standard resume sections",
        "score": section_score,
        "max_score": 20,
    })

    # 3. Job-description keyword alignment: 35 points
    resume_keywords = extract_keywords(resume_text)
    job_keywords = extract_keywords(job_description)

    if job_keywords:
        matched_keywords = resume_keywords.intersection(job_keywords)
        keyword_match_ratio = len(matched_keywords) / len(job_keywords)
    else:
        matched_keywords = set()
        keyword_match_ratio = 0.0

    keyword_score = round(min(keyword_match_ratio, 1.0) * 35)
    score += keyword_score

    breakdown.append({
        "category": "Job-description keyword alignment",
        "score": keyword_score,
        "max_score": 35,
    })

    # 4. Measurable achievements: 15 points
    number_patterns = re.findall(
        r"\b\d+(?:\.\d+)?\s*(?:%|percent|k|m|million|thousand|x)?\b",
        resume_text.lower(),
    )

    achievement_score = min(15, len(number_patterns) * 3)
    score += achievement_score

    breakdown.append({
        "category": "Measurable achievements",
        "score": achievement_score,
        "max_score": 15,
    })

    # 5. Basic formatting/readability signals: 15 points
    formatting_score = 0

    if len(resume_text.splitlines()) >= 10:
        formatting_score += 5

    if not re.search(r"[^\x00-\x7F]", resume_text):
        formatting_score += 3

    if resume_text.count("\n") >= 5:
        formatting_score += 3

    if not re.search(r"(.)\1{5,}", resume_text):
        formatting_score += 2

    if "@" in resume_text:
        formatting_score += 2

    formatting_score = min(formatting_score, 15)
    score += formatting_score

    breakdown.append({
        "category": "Basic formatting and readability",
        "score": formatting_score,
        "max_score": 15,
    })

    return {
        "score": max(0, min(100, score)),
        "word_count": word_count,
        "detected_sections": detected_sections,
        "matched_keywords": sorted(matched_keywords),
        "job_keywords_count": len(job_keywords),
        "breakdown": breakdown,
    }


# =========================================================
# Gemini integration
# =========================================================

def get_gemini_client() -> genai.Client:
    """Create a Gemini client using Streamlit secrets."""
    api_key = st.secrets.get("GEMINI_API_KEY")

    if not api_key:
        raise ValueError(
            "GEMINI_API_KEY is missing. Add it in Streamlit Secrets or "
            "in .streamlit/secrets.toml for local testing."
        )

    return genai.Client(api_key=api_key)


def clean_json_response(response_text: str) -> str:
    """Remove accidental Markdown code fences around JSON."""
    cleaned = response_text.strip()

    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned)

    return cleaned.strip()


def safe_list(value: Any) -> List[Any]:
    """Return a list or an empty list."""
    return value if isinstance(value, list) else []


def analyze_with_gemini(
    resume_text: str,
    job_description: str,
    basic_score: Dict[str, Any],
) -> Dict[str, Any]:
    """Analyze resume content and job alignment using Gemini."""
    client = get_gemini_client()

    prompt = f"""
You are an expert resume reviewer and ATS optimization consultant.

Analyze the resume against the job description.

Rules:
1. This is an estimated analysis, not the official score of any ATS platform.
2. Return an AI content and alignment score from 0 to 100.
3. Never invent experience, education, employers, certifications, skills,
   achievements, dates, or metrics.
4. Only recommend adding a keyword if it is relevant to the job description.
5. Clearly tell the user to verify any missing keyword before adding it.
6. Rewrite bullets using only facts already present in the resume.
7. Be specific, practical, and concise.
8. Return ONLY valid JSON. Do not return Markdown.

Return exactly this JSON structure:

{{
  "ai_score": 0,
  "summary": "Short overall assessment",
  "strengths": ["...", "..."],
  "weaknesses": ["...", "..."],
  "missing_keywords": ["...", "..."],
  "formatting_issues": ["...", "..."],
  "content_improvements": ["...", "..."],
  "rewritten_bullets": [
    {{
      "original": "Existing resume bullet",
      "improved": "Improved version using only existing facts",
      "reason": "Why this is better"
    }}
  ],
  "recommended_sections": ["...", "..."],
  "final_recommendation": "Most important next step"
}}

Resume:
--- BEGIN RESUME ---
{resume_text[:MAX_RESUME_CHARS]}
--- END RESUME ---

Job Description:
--- BEGIN JOB DESCRIPTION ---
{job_description[:MAX_JOB_DESCRIPTION_CHARS]}
--- END JOB DESCRIPTION ---

Rule-based analysis:
{json.dumps(basic_score, indent=2)}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )

    response_text = getattr(response, "text", None)

    if not response_text:
        raise ValueError("Gemini returned an empty response.")

    try:
        result = json.loads(clean_json_response(response_text))
    except json.JSONDecodeError as exc:
        raise ValueError(
            "Gemini returned an invalid JSON response. Please try again."
        ) from exc

    if not isinstance(result, dict):
        raise ValueError("Gemini returned an unexpected response format.")

    try:
        ai_score = float(result.get("ai_score", 0))
    except (TypeError, ValueError):
        ai_score = 0

    result["ai_score"] = max(0, min(100, ai_score))

    list_fields = [
        "strengths",
        "weaknesses",
        "missing_keywords",
        "formatting_issues",
        "content_improvements",
        "recommended_sections",
    ]

    for field in list_fields:
        result[field] = safe_list(result.get(field))

    rewritten = safe_list(result.get("rewritten_bullets"))
    result["rewritten_bullets"] = [
        item for item in rewritten if isinstance(item, dict)
    ]

    result["summary"] = str(result.get("summary", ""))
    result["final_recommendation"] = str(
        result.get("final_recommendation", "")
    )

    return result


# =========================================================
# UI helpers
# =========================================================

def score_label(score: float) -> str:
    if score >= 80:
        return "Strong"
    if score >= 65:
        return "Good, but improvable"
    if score >= 50:
        return "Needs improvement"
    return "Weak alignment"


def display_list(
    items: List[Any],
    empty_message: str = "None identified.",
) -> None:
    if not items:
        st.info(empty_message)
        return

    for item in items:
        if isinstance(item, dict):
            st.markdown(f"**{item.get('original', 'Suggestion')}**")
            st.write(item.get("improved", ""))
            if item.get("reason"):
                st.caption(item["reason"])
        else:
            st.markdown(f"- {item}")


# =========================================================
# Main Streamlit app
# =========================================================

st.title("📄 Resume ATS Analyzer")
st.caption(
    "Upload your resume and receive an estimated ATS compatibility score, "
    "keyword analysis, and AI-powered improvement suggestions."
)

st.warning(
    "Important: This is an estimated score. It is not an official score "
    "from Workday, Greenhouse, Lever, LinkedIn, or another ATS platform."
)

with st.sidebar:
    st.header("How it works")
    st.markdown(
        """
        1. Upload a PDF or DOCX resume.
        2. Paste the target job description.
        3. Resume text is extracted locally.
        4. A transparent rule-based score is calculated.
        5. Gemini reviews content and job alignment.
        6. The final report combines both analyses.
        """
    )

    st.divider()
    st.write("**Gemini model:**")
    st.code(MODEL_NAME)

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"],
    help="Maximum file size: 10 MB.",
)

job_description = st.text_area(
    "Paste the target job description",
    height=260,
    placeholder=(
        "Paste the complete job description here. "
        "For example: We are looking for a Python Developer with experience "
        "in Django, REST APIs, SQL, Git, and cloud deployment..."
    ),
)

analyze_button = st.button(
    "🔍 Analyze Resume",
    type="primary",
    use_container_width=True,
)

if analyze_button:
    if uploaded_file is None:
        st.error("Please upload a PDF or DOCX resume.")
        st.stop()

    if not job_description.strip():
        st.error(
            "Please paste a job description for a meaningful ATS analysis."
        )
        st.stop()

    with st.spinner("Extracting resume text..."):
        try:
            resume_text = extract_resume_text(uploaded_file)
        except Exception as exc:
            st.error(f"Could not read the resume: {exc}")
            st.stop()

    with st.spinner("Calculating rule-based ATS score..."):
        basic_score = calculate_basic_ats_score(
            resume_text,
            job_description,
        )

    with st.spinner("Gemini is reviewing your resume..."):
        try:
            ai_result = analyze_with_gemini(
                resume_text,
                job_description,
                basic_score,
            )
        except Exception as exc:
            st.error(f"Gemini analysis failed: {exc}")
            st.stop()

    # Weighted final score:
    # 40% transparent rule-based score
    # 60% Gemini content and alignment score
    final_score = round(
        basic_score["score"] * 0.40
        + ai_result["ai_score"] * 0.60
    )

    st.session_state["analysis"] = {
        "resume_text": resume_text,
        "basic_score": basic_score,
        "ai_result": ai_result,
        "final_score": final_score,
    }

if "analysis" in st.session_state:
    analysis = st.session_state["analysis"]
    basic_score = analysis["basic_score"]
    ai_result = analysis["ai_result"]
    final_score = analysis["final_score"]

    st.divider()
    st.header("Your Resume Analysis")

    col1, col2, col3 = st.columns(3)

    with col1:
        st.metric("Estimated ATS Score", f"{final_score}/100")

    with col2:
        st.metric(
            "Rule-Based Score",
            f"{basic_score['score']}/100",
        )

    with col3:
        st.metric(
            "AI Content Score",
            f"{ai_result['ai_score']:.0f}/100",
        )

    st.progress(final_score / 100)
    st.success(f"Overall assessment: **{score_label(final_score)}**")

    st.info(
        "Final score = 40% rule-based analysis + "
        "60% AI content and job-alignment analysis."
    )

    tab1, tab2, tab3, tab4, tab5 = st.tabs(
        [
            "📊 Score Breakdown",
            "✅ Strengths & Weaknesses",
            "🔑 Keywords",
            "✍️ Rewrite Suggestions",
            "📋 Resume Text",
        ]
    )

    with tab1:
        st.subheader("Rule-Based Score Breakdown")

        for item in basic_score["breakdown"]:
            st.write(
                f"**{item['category']}**: "
                f"{item['score']}/{item['max_score']}"
            )
            st.progress(item["score"] / item["max_score"])

        st.write(f"**Resume word count:** {basic_score['word_count']}")

        st.subheader("Detected Resume Sections")
        detected_sections = basic_score.get("detected_sections", [])
        if detected_sections:
            st.write(", ".join(detected_sections))
        else:
            st.info("No standard sections were detected.")

        st.subheader("AI Summary")
        st.write(
            ai_result.get(
                "summary",
                "No summary available.",
            )
        )

        st.subheader("Final Recommendation")
        st.write(
            ai_result.get(
                "final_recommendation",
                "No recommendation available.",
            )
        )

    with tab2:
        left, right = st.columns(2)

        with left:
            st.subheader("Strengths")
            display_list(ai_result.get("strengths", []))

        with right:
            st.subheader("Weaknesses")
            display_list(ai_result.get("weaknesses", []))

        st.subheader("Formatting Issues")
        display_list(ai_result.get("formatting_issues", []))

        st.subheader("Content Improvements")
        display_list(ai_result.get("content_improvements", []))

    with tab3:
        st.subheader("Keyword Alignment")

        matched_keywords = basic_score.get("matched_keywords", [])
        missing_keywords = ai_result.get("missing_keywords", [])

        st.write(f"**Matched keyword count:** {len(matched_keywords)}")

        if matched_keywords:
            st.write(", ".join(matched_keywords))
        else:
            st.info("No significant matching keywords were detected.")

        st.subheader("Potential Missing Keywords")
        st.caption(
            "Verify these against your actual experience before adding them."
        )
        display_list(
            missing_keywords,
            "No major missing keywords were identified.",
        )

    with tab4:
        st.subheader("Rewritten Bullet Suggestions")

        rewritten_bullets = ai_result.get("rewritten_bullets", [])

        if not rewritten_bullets:
            st.info("No bullet rewrites were generated.")
        else:
            for index, bullet in enumerate(rewritten_bullets, start=1):
                st.markdown(f"### Suggestion {index}")
                st.markdown("**Original**")
                st.write(bullet.get("original", ""))
                st.markdown("**Improved**")
                st.write(bullet.get("improved", ""))
                st.caption(bullet.get("reason", ""))
                st.divider()

        st.subheader("Recommended Resume Sections")
        display_list(ai_result.get("recommended_sections", []))

    with tab5:
        st.subheader("Extracted Resume Text")
        st.caption(
            "This is the text sent to Gemini. Review it if your PDF "
            "contains unusual formatting."
        )
        st.text_area(
            "Resume text",
            value=analysis["resume_text"],
            height=500,
            disabled=True,
        )

    st.divider()
    st.caption(
        "Privacy note: Resume text is sent to Gemini for analysis. "
        "Do not upload sensitive documents unless you are comfortable "
        "with the API provider's data-handling policies."
    )
